# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def set_cell(cell, text, bold_first_line=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for idx, line in enumerate(text.split("\n")):
        if idx > 0:
            p = cell.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9.5)
        if bold_first_line and idx == 0:
            run.bold = True
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_cell(doc, table_idx, row, col, text, bold_first_line=False):
    set_cell(doc.tables[table_idx].cell(row, col), text, bold_first_line)


def main():
    src_candidates = list((Path.home() / "Downloads").glob("*.docx"))
    src = src_candidates[1]
    out = Path("MLAgent_Report_DumDum_completed.docx").resolve()

    doc = Document(src)
    try:
        style = doc.styles["Normal"]
        style.font.name = "Malgun Gothic"
        style.font.size = Pt(10)
    except Exception:
        pass

    set_cell(doc.tables[0].cell(0, 0), "2026년 ML-Agent 프로젝트 보고서\nDumDum - 고민 해결 사고 과정 훈련 앱", True)
    doc.tables[0].cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell(doc.tables[1].cell(0, 0), "프로젝트 파일 링크: 제출 전 프로젝트 ZIP 업로드 후 링크 삽입\n※ 현재 작업 폴더: C:/Users/sobin/DumDum")
    set_cell(doc.tables[2].cell(0, 0), "빌드본 파일 링크: 제출 전 Android 빌드본 ZIP 업로드 후 링크 삽입\n※ 빌드 시 알림 테스트를 위해 알림 지연 시간은 1분으로 설정함")

    activity = [
        ("1회차", "프로젝트 주제 선정 및 기획 정리\n- 막연한 걱정을 문장화하고 단계별 사고 과정으로 정리하는 앱으로 방향 설정\n- 핵심 루프를 고민 입력, 확률 예상, 최악 상황 대비, 행동 정리, 감정 선택, 결과 태그, 회고 학습으로 설계"),
        ("2회차", "기본 앱 UI와 화이트보드 기능 구현\n- 고민 카드를 저장하고 화이트보드에 누적하는 구조 구현\n- 고민 상세 보기, 실제 결과 태그, 감정 상태 기록, 결과 통계 표시 기능 추가"),
        ("3회차", "회고 학습 기능 및 알림 구조 구현\n- 결과 태그 후 처음 예상 확률과 실제 결과를 비교하는 회고 문장 입력 UI 추가\n- 카드 저장 및 결과 회고 후 사용자 상태에 맞는 지원 알림을 예약하는 서비스 구현"),
        ("4회차", "ML-Agent 학습 환경 구성\n- 사용자 상태를 관측값으로 만들고 알림 종류를 행동으로 선택하는 DumDumSupportAgent 구현\n- 가상 사용자 시나리오와 보상 규칙을 가진 학습 씬 제작"),
        ("5회차", "학습 실행 및 빌드 테스트 준비\n- Python ML-Agents 환경을 구성하고 100,000 step 학습 버튼/런처 제작\n- 학습 중 연결 문제와 Unity 도메인 리로드 문제를 확인하고 대기 시간, 로그 저장, 실행 순서를 개선"),
    ]
    for i, (round_name, desc) in enumerate(activity, start=1):
        set_table_cell(doc, 3, i, 0, round_name)
        set_table_cell(doc, 3, i, 1, desc)

    set_table_cell(doc, 4, 0, 1, "DumDum - 고민 해결 사고 과정 훈련 앱")
    set_table_cell(doc, 4, 1, 1, "DumDum은 사용자가 고민을 머릿속에 막연하게 가지고 있다가 불안이 커지는 상황을 줄이기 위한 앱이다. 사용자는 걱정 내용을 문장으로 입력하고, 그 일이 실제로 일어날 확률, 최악의 경우 대처 방법, 지금 할 수 있는 행동, 현재 감정 상태를 단계별로 정리한다. 이후 시간이 지난 뒤 실제 결과를 태그하고 회고 문장을 남기며, 자신의 걱정이 얼마나 현실과 달랐는지 시각적으로 학습한다.")
    set_table_cell(doc, 4, 2, 1, "1. 시작 화면에서 새 고민 정리를 선택한다.\n2. 고민 내용을 입력하고 예상 확률을 슬라이더로 조정한다.\n3. 최악의 경우 할 수 있는 행동과 지금 할 수 있는 행동을 작성한다.\n4. 현재 감정 상태를 선택하면 고민 카드가 화이트보드에 저장된다.\n5. 저장된 카드를 선택해 실제 결과를 태그하고 회고 문장을 남긴다.\n6. 앱은 조건에 따라 결과 확인 또는 회고 알림을 예약한다.")
    set_table_cell(doc, 4, 3, 1, "ML-Agent는 앱 본편의 직접 조작 캐릭터가 아니라, 사용자 상태에 맞는 지원 알림을 선택하는 학습 에이전트로 적용하였다. 고민 카드 수, 태그된 결과 수, 일어나지 않은 걱정 비율, 오늘 작성한 고민 수, 회고 문장 수, 최근 감정 상태 등을 관측값으로 사용하고, 에이전트는 “알림 없음”, “과거 회고 문장 알림”, “걱정 실제 발생률 알림”, “하루 마무리 권유 알림” 중 하나를 선택하도록 학습한다.")

    core = """1. 화이트보드 기반 고민 카드 시스템
- 사용자가 정리한 고민을 카드 형태로 저장하고, 감정 상태와 행동 계획을 함께 표시한다.
- 단순 기록이 아니라 과거의 걱정이 어떻게 지나갔는지를 보는 증거 보드 역할을 하도록 구성하였다.

2. 단계별 사고 정리 플로우
- 고민 입력 → 예상 확률 선택 → 최악 상황 대비 → 현재 할 수 있는 행동 작성 → 감정 선택 순서로 진행된다.
- 막연한 불안을 구체적인 문장과 행동 계획으로 바꾸는 것이 핵심이다.

3. 실제 결과 태그와 회고 학습
- 저장된 고민 카드에 대해 “일어나지 않음”, “일부만 일어남”, “실제로 일어남” 결과 태그를 남길 수 있다.
- 태그 후에는 처음 예상 확률과 실제 결과를 비교하는 회고 문장을 작성하여 예측 오류를 인식하도록 만들었다.

4. 결과 통계 및 시각적 피드백
- 태그된 카드의 비율을 바탕으로 실제로 일어나지 않은 걱정의 비율을 보여준다.
- 사용자가 “내 걱정은 생각보다 자주 현실이 되지 않았다”는 패턴을 스스로 확인할 수 있게 한다.

5. 모바일 알림 기능
- 고민 저장 후 결과 확인 알림을 예약하고, 회고 문장이나 걱정 통계가 충분할 때 지원 알림을 예약한다.
- Android Mobile Notifications 패키지를 사용하여 실제 기기에서 알림을 받을 수 있도록 구현하였다.

6. ML-Agent 학습 씬
- 실제 앱 데이터와 비슷한 가상 사용자 상태를 생성하고, 여러 개의 에이전트가 반복 학습하도록 구성하였다.
- 100,000 step 학습을 목표로 하는 버튼과 트레이너 실행 스크립트를 제작하였다."""
    set_cell(doc.tables[5].cell(0, 0), core)

    set_table_cell(doc, 6, 1, 1, "DumDumSupportAgent는 Initialize()에서 학습 환경인 DumDumSupportTrainingEnvironment를 찾거나 연결한다. 학습 씬에는 여러 개의 에이전트 오브젝트가 배치되어 있고, 각 에이전트는 같은 Behavior Name(DumDumSupportIntervention)을 사용한다.")
    set_table_cell(doc, 6, 2, 1, "OnEpisodeBegin()에서 학습 환경이 새로운 가상 사용자 시나리오를 만든다. 시나리오는 크게 하루 마무리 권유가 필요한 상황, 회고 문장 알림이 필요한 상황, 결과 통계 알림이 필요한 상황, 알림이 필요 없는 상황으로 나뉜다.")
    set_table_cell(doc, 6, 3, 1, "CollectObservations()에서 총 12개의 관측값을 VectorSensor에 추가한다. 관측값은 전체 고민 카드 수, 태그된 카드 수, 미태그 카드 수, 일어나지 않은 걱정 비율, 실제로 일어난 걱정 비율, 오늘 작성한 고민 수, 회고 문장 수, 최근 예상 확률, 최근 감정 상태, 회고 문장 존재 여부, 심란한 감정 여부, 결과 데이터 존재 여부이다. 숫자 값은 0~1 범위로 정규화하였다.")
    set_table_cell(doc, 6, 4, 1, "OnActionReceived()에서 에이전트는 Discrete Action 하나를 선택한다. 행동 0은 알림 없음, 1은 과거 회고 문장 알림, 2는 걱정 실제 발생률 알림, 3은 하루 마무리 권유 알림이다. 선택한 행동은 SupportInterventionType으로 변환되고, 학습 환경의 Evaluate() 함수가 보상을 계산한다.")
    set_table_cell(doc, 6, 5, 1, "Heuristic()에는 규칙 기반 선택 로직을 넣었다. 오늘 고민이 2개 이상이고 감정이 심란하면 하루 마무리 권유, 회고 문장이 있으면 회고 문장 알림, 태그된 카드가 3개 이상이고 일어나지 않은 비율이 50% 이상이면 결과 통계 알림을 선택한다. 이 로직은 학습 전 동작 확인과 비교 기준으로 사용하였다.")
    set_table_cell(doc, 6, 6, 1, "정답 행동을 선택하면 +1 보상을 주고, 틀린 행동을 선택하면 -0.35 보상을 준다. 특히 알림이 필요 없는 상황에서 불필요한 알림을 보내면 추가 감점을 주어 과도한 알림을 줄이도록 설계하였다. 알림이 필요 없는 상황에서 알림 없음 행동을 선택하면 0.65 보상을 주었다.")

    learning = """학습은 Unity ML-Agents 4.0.3 패키지와 Python mlagents 1.1.0 환경을 사용하여 진행하였다. Behavior Name은 DumDumSupportIntervention으로 설정했고, PPO 알고리즘을 사용하였다. 학습 설정은 batch_size 128, buffer_size 2048, learning_rate 0.0003, hidden_units 64, num_layers 2, max_steps 100000으로 구성하였다.

처음에는 Python 3.10 환경과 mlagents 패키지 버전이 맞지 않아 학습 실행이 되지 않는 문제가 있었다. 이를 해결하기 위해 프로젝트 로컬 Python 3.10 런타임과 학습 실행 스크립트를 준비했고, 환경 점검 스크립트를 통해 ml-agents, ml-agents-envs, torch가 정상적으로 import되는지 확인하였다.

다음 문제는 Unity Editor와 Python trainer의 연결 순서였다. ML-Agents는 trainer가 먼저 Listening 상태가 된 뒤 Unity Play Mode에 들어가야 안정적으로 연결되는데, 처음에는 씬 안 버튼을 Play 상태에서 눌러 trainer가 늦게 실행되면서 연결이 실패했다. 이를 해결하기 위해 trainer를 먼저 실행하고 일정 시간 뒤 자동으로 Play Mode에 들어가는 런처를 추가하였다.

학습 중에는 Unity가 스크립트 컴파일 또는 도메인 리로드를 수행하면 trainer와의 통신이 끊기는 문제가 있었다. 이 문제를 확인하기 위해 PowerShell 창이 바로 닫히지 않게 하고, MLRuns/_trainer_logs 폴더에 trainer 로그를 저장하도록 수정하였다. 이후 학습 중에는 코드 수정이나 씬 변경을 하지 않는 방식으로 안정성을 높였다.

학습 로그에서 DumDumSupportIntervention이 Unity 환경과 연결되고, Step 로그가 출력되는 것을 확인하였다. 초반 Mean Reward는 낮게 시작했지만, 에이전트가 여러 사용자 상태에서 적절한 알림 행동을 선택하도록 반복 학습하는 구조가 정상적으로 동작함을 확인하였다. 최종적으로 학습 씬과 앱 본편의 규칙 기반 알림 정책을 연결하여, ML-Agent를 활용한 알림 선택 정책 설계와 학습 과정을 프로젝트에 포함하였다."""
    set_cell(doc.tables[7].cell(0, 0), learning)

    reflection = """이번 프로젝트를 통해 ML-Agent는 단순히 적을 따라오게 하거나 장애물을 피하게 하는 용도뿐 아니라, 사용자의 상태를 보고 적절한 개입을 선택하는 방식으로도 활용할 수 있다는 점을 배웠다. 처음에는 “고민 정리 앱에 ML-Agent를 어떻게 넣을 수 있을까?”가 가장 어려웠지만, 앱의 알림 선택 문제를 하나의 의사결정 문제로 바꾸면서 적용 방향이 명확해졌다.

기존 게임 인공지능인 FSM이나 Behavior Tree는 개발자가 조건과 행동을 직접 설계한다. 예를 들어 “체력이 낮으면 도망간다”, “플레이어가 가까우면 공격한다”처럼 규칙이 분명하다. 반면 ML-Agent는 관측값, 행동, 보상 구조를 설계하면 반복 학습을 통해 어떤 행동이 더 좋은지 스스로 정책을 찾아간다. 이 점에서 FSM/BT는 예측 가능하고 디버깅이 쉽지만, 상황이 복잡해질수록 규칙이 많아지는 단점이 있다. ML-Agent는 학습 과정이 필요하고 결과를 해석하기 어렵지만, 다양한 상태를 반복 경험하면서 정책을 개선할 수 있다는 장점이 있다.

이번 앱에서는 사용자의 고민 수, 감정 상태, 회고 여부, 실제 발생률 같은 데이터를 관측값으로 사용하였다. 이 데이터는 게임에서의 체력, 거리, 탄약 수 같은 상태값과 비슷한 역할을 한다. 에이전트가 선택하는 행동은 알림 없음, 회고 문장 알림, 통계 알림, 하루 마무리 권유였고, 보상은 상황에 맞는 선택을 했는지에 따라 주었다. 이를 통해 앱의 기능도 게임 AI처럼 상태-행동-보상 구조로 바꿀 수 있다는 점을 경험하였다.

아쉬운 점도 있었다. 실제 사용자의 데이터를 바로 학습에 사용한 것이 아니라, 규칙 기반으로 만든 가상 시나리오에서 학습했기 때문에 완전히 개인화된 AI라고 보기는 어렵다. 또한 학습된 모델을 앱 본편에 직접 탑재하기보다는, 학습 설계와 같은 기준의 알림 정책을 본편에 연결한 형태라서 추후에는 ONNX 모델을 앱에 적용하고 실제 추론 결과로 알림을 고르는 구조까지 확장해보고 싶다.

그래도 이번 프로젝트는 ML-Agent를 단순 게임 캐릭터가 아닌 사용자 경험 개선 도구로 적용해본 점에서 의미가 있었다. 앞으로 게임 업계에서도 ML은 적 AI, NPC 행동, 난이도 조절뿐 아니라 유저 상태 분석, 튜토리얼 추천, 개인화된 도움 제공 같은 방향으로 더 많이 활용될 수 있을 것이라고 생각한다. 중요한 것은 AI를 넣는 것 자체가 아니라, 어떤 문제를 학습 가능한 형태로 바꾸고 어떤 보상을 줄 것인지 설계하는 능력이라는 것을 느꼈다."""
    set_cell(doc.tables[8].cell(0, 0), reflection)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    for r in p.runs:
                        r.font.name = "Malgun Gothic"
                        if r.font.size is None:
                            r.font.size = Pt(9.5)

    for table_idx in [3, 4, 6]:
        for cell in doc.tables[table_idx].rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
